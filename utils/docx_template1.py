"""
NEAB DOCX Template Generator (Python version)
Skapar professionella Word-rapporter för säkerhetsbedömningar
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
import logging

logger = logging.getLogger(__name__)

# NEAB Färgschema (samma som PDF)
class NEABColors:
    PRIMARY = RGBColor(30, 58, 138)      # Mörk blå
    SECONDARY = RGBColor(59, 130, 246)   # Ljus blå
    ACCENT = RGBColor(16, 185, 129)      # Grön
    SUCCESS = RGBColor(16, 185, 129)
    WARNING = RGBColor(245, 158, 11)
    DANGER = RGBColor(239, 68, 68)
    NEUTRAL = RGBColor(107, 114, 128)
    LIGHT_GRAY = RGBColor(243, 244, 246)
    DARK_GRAY = RGBColor(55, 65, 81)


class NEABDocxTemplate:
    """Professionell DOCX-template för NEAB-rapporter"""

    def __init__(self):
        self.doc = Document()
        self._setup_styles()
        self._setup_sections()

    def _setup_styles(self):
        """Skapar anpassade stilar"""
        styles = self.doc.styles

        # Heading 1 style
        if 'NEAB Heading 1' not in [s.name for s in styles]:
            h1_style = styles.add_style('NEAB Heading 1', WD_STYLE_TYPE.PARAGRAPH)
            h1_style.base_style = styles['Heading 1']
            h1_font = h1_style.font
            h1_font.name = 'Arial'
            h1_font.size = Pt(18)
            h1_font.bold = True
            h1_font.color.rgb = NEABColors.PRIMARY
            h1_style.paragraph_format.space_before = Pt(24)
            h1_style.paragraph_format.space_after = Pt(12)

        # Heading 2 style
        if 'NEAB Heading 2' not in [s.name for s in styles]:
            h2_style = styles.add_style('NEAB Heading 2', WD_STYLE_TYPE.PARAGRAPH)
            h2_style.base_style = styles['Heading 2']
            h2_font = h2_style.font
            h2_font.name = 'Arial'
            h2_font.size = Pt(14)
            h2_font.bold = True
            h2_font.color.rgb = NEABColors.SECONDARY
            h2_style.paragraph_format.space_before = Pt(18)
            h2_style.paragraph_format.space_after = Pt(10)

        # Heading 3 style
        if 'NEAB Heading 3' not in [s.name for s in styles]:
            h3_style = styles.add_style('NEAB Heading 3', WD_STYLE_TYPE.PARAGRAPH)
            h3_style.base_style = styles['Heading 3']
            h3_font = h3_style.font
            h3_font.name = 'Arial'
            h3_font.size = Pt(12)
            h3_font.bold = True
            h3_font.color.rgb = NEABColors.DARK_GRAY
            h3_style.paragraph_format.space_before = Pt(12)
            h3_style.paragraph_format.space_after = Pt(8)

        # Body style
        if 'NEAB Body' not in [s.name for s in styles]:
            body_style = styles.add_style('NEAB Body', WD_STYLE_TYPE.PARAGRAPH)
            body_style.base_style = styles['Normal']
            body_font = body_style.font
            body_font.name = 'Arial'
            body_font.size = Pt(11)
            body_font.color.rgb = NEABColors.DARK_GRAY
            body_style.paragraph_format.space_after = Pt(8)
            body_style.paragraph_format.line_spacing = 1.15

    def _setup_sections(self):
        """Konfigurerar sektioner med marginaler"""
        section = self.doc.sections[0]
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    def add_cover_page(self, data):
        """Skapar försättssida"""
        # Logo info (om tillgänglig)
        if data.get('logo_path'):
            p = self.doc.add_paragraph('[LOGO]')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.runs[0]
            run.font.bold = True
            run.font.color.rgb = NEABColors.NEUTRAL
            self.doc.add_paragraph()

        # Mellanrum
        for _ in range(3):
            self.doc.add_paragraph()

        # Huvudtitel
        p = self.doc.add_paragraph(data.get('title', 'NEAB Säkerhetsrapport'))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.font.name = 'Arial'
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = NEABColors.PRIMARY

        self.doc.add_paragraph()

        # Undertitel
        p = self.doc.add_paragraph(data.get('subtitle', 'NIS2 och regulatorisk efterlevnad'))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.font.name = 'Arial'
        run.font.size = Pt(16)
        run.font.color.rgb = NEABColors.SECONDARY

        # Mellanrum
        for _ in range(2):
            self.doc.add_paragraph()

        # Organisation
        p = self.doc.add_paragraph(data.get('organization', 'Nordlunda Energi AB'))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.font.name = 'Arial'
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = NEABColors.PRIMARY

        self.doc.add_paragraph()

        # Datum
        date_str = data.get('date', '2025-12-15')
        p = self.doc.add_paragraph(f'Datum: {date_str}')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0]
        run.font.name = 'Arial'
        run.font.size = Pt(16)
        run.font.color.rgb = NEABColors.SECONDARY

        # Sidbrytning
        self.doc.add_page_break()

    def add_executive_summary(self, data):
        """Skapar sammanfattning för ledningen"""
        # Huvudrubrik
        p = self.doc.add_paragraph('Sammanfattning för ledningsnivå')
        p.style = 'NEAB Heading 1'

        # Övergripande bedömning
        p = self.doc.add_paragraph('Övergripande bedömning:')
        p.style = 'NEAB Heading 2'

        compliance_rate = int(round(data.get('compliance_rate', 0)))
        assessment = (f"NEAB har en {compliance_rate}% regelefterlevnad med "
                     f"{data.get('very_high_risks', 0)} mycket höga och "
                     f"{data.get('high_risks', 0)} höga risker identifierade.")

        p = self.doc.add_paragraph(assessment)
        p.style = 'NEAB Body'

        # Viktigaste riskområden
        p = self.doc.add_paragraph('Viktigaste riskområden:')
        p.style = 'NEAB Heading 2'

        for risk in data.get('top_risks', [])[:5]:
            p = self.doc.add_paragraph(risk, style='List Bullet')
            p.paragraph_format.space_after = Pt(4)

        self.doc.add_paragraph()

        # Regulatoriska brister
        p = self.doc.add_paragraph('Regulatoriska brister:')
        p.style = 'NEAB Heading 2'

        for gap in data.get('regulatory_gaps', [])[:5]:
            p = self.doc.add_paragraph(gap, style='List Bullet')
            p.paragraph_format.space_after = Pt(4)

        self.doc.add_paragraph()

        # Föreslagna åtgärder
        p = self.doc.add_paragraph('Föreslagna prioriterade åtgärder:')
        p.style = 'NEAB Heading 2'

        for action in data.get('priority_actions', [])[:5]:
            p = self.doc.add_paragraph(action, style='List Number')
            p.paragraph_format.space_after = Pt(4)

        # Sidbrytning
        self.doc.add_page_break()

    def add_organization_section(self, data):
        """Skapar organisationssektion"""
        p = self.doc.add_paragraph('Organisationen och kritiska processer')
        p.style = 'NEAB Heading 1'

        # Kärnverksamhet
        p = self.doc.add_paragraph('Kärnverksamhet:')
        p.style = 'NEAB Heading 2'

        core_business = data.get('core_business',
                                'Nordlunda Energi AB är ett regionalt energibolag som driver '
                                'elnätsverksamhet (DSO) och fjärrvärme. Bolaget betjänar ca 85000 '
                                'elnätskunder och 12000 fjärrvärmekunder.')
        p = self.doc.add_paragraph(core_business)
        p.style = 'NEAB Body'

        # Kritiska system
        p = self.doc.add_paragraph('Kritiska system:')
        p.style = 'NEAB Heading 2'

        critical_systems = data.get('critical_systems',
                                   'SCADA/DCS för styrning, AMI för mätdata, OMS för avbrottshantering, '
                                   'samt kund- och faktureringssystem.')
        p = self.doc.add_paragraph(critical_systems)
        p.style = 'NEAB Body'

        # Säkerhetsstyrning
        p = self.doc.add_paragraph('Säkerhetsstyrning:')
        p.style = 'NEAB Heading 2'

        security_gov = data.get('security_governance',
                               'CISO med stab, OT-säkerhetsansvarig, samt SOC via MSSP '
                               'med 24/7-övervakning.')
        p = self.doc.add_paragraph(security_gov)
        p.style = 'NEAB Body'

        # Sidbrytning
        self.doc.add_page_break()

    def add_regulatory_summary(self, data):
        """Skapar regulatorisk sammanfattning"""
        p = self.doc.add_paragraph('Regulatorisk kartläggning')
        p.style = 'NEAB Heading 1'

        total_reqs = data.get('total_requirements', 9)
        fulfilled = data.get('fulfilled', 2)
        partial = data.get('partial', 6)
        not_fulfilled = data.get('not_fulfilled', 1)
        compliance_rate = int(round(data.get('compliance_rate', 55.6)))

        summary = (f"Nordlunda Energi AB har en otillräcklig efterlevnadsnivå ({compliance_rate}%) "
                  f"av tillämpliga regelverk. Av {total_reqs} identifierade krav: "
                  f"{fulfilled} fullt uppfyllda, {partial} delvis uppfyllda, {not_fulfilled} ej uppfyllda.")

        p = self.doc.add_paragraph(summary)
        p.style = 'NEAB Body'

        # Prioriterade områden
        p = self.doc.add_paragraph('Prioriterade områden för förbättring:')
        p.style = 'NEAB Heading 2'

        priorities = data.get('priority_areas',
                            ['Styrelseutbildning (NIS2)',
                             'DPIA för mätdata (GDPR)',
                             'Formalisering av riskhanteringsprocesser'])
        for priority in priorities:
            p = self.doc.add_paragraph(priority, style='List Bullet')
            p.paragraph_format.space_after = Pt(4)

        # Sidbrytning
        self.doc.add_page_break()

    def add_risk_summary(self, data):
        """Skapar risksammanfattning"""
        p = self.doc.add_paragraph('Sammanfattning av riskanalys')
        p.style = 'NEAB Heading 1'

        risks = data.get('risk_distribution', {})
        total = sum(risks.values())

        summary = (f"Riskworkshop identifierade {total} risker: "
                  f"{risks.get('Mycket hög', 0)} mycket höga risker (kräver omedelbar åtgärd), "
                  f"{risks.get('Hög', 0)} höga risker (prioriterad hantering), "
                  f"{risks.get('Medel', 0)} medelhöga risker (planerad hantering), "
                  f"{risks.get('Låg', 0)} låga risker (övervakning).")

        p = self.doc.add_paragraph(summary)
        p.style = 'NEAB Body'

        # Top 3 risker
        p = self.doc.add_paragraph('Top 3 risker:')
        p.style = 'NEAB Heading 2'

        for risk in data.get('top_risks', [])[:3]:
            p = self.doc.add_paragraph(risk, style='List Number')
            p.paragraph_format.space_after = Pt(4)

        # Info om heat map
        if data.get('heatmap_path'):
            self.doc.add_paragraph()
            p = self.doc.add_paragraph('[Heat map/riskmatris finns i PDF-versionen]')
            run = p.runs[0]
            run.font.italic = True
            run.font.color.rgb = NEABColors.NEUTRAL

        # Sidbrytning
        self.doc.add_page_break()

    def add_mitigation_summary(self, data):
        """Skapar mitigeringssammanfattning"""
        p = self.doc.add_paragraph('Mitigeringsstrategier och rekommenderade åtgärder')
        p.style = 'NEAB Heading 1'

        priority_actions = data.get('priority_actions', [])
        num_actions = len(priority_actions)

        p = self.doc.add_paragraph(f'Identifierade {num_actions} prioriterade åtgärder')
        p.style = 'NEAB Body'

        # Lista åtgärder
        if priority_actions and priority_actions[0] != 'Okänd åtgärd':
            p = self.doc.add_paragraph('Top prioriterade åtgärder:')
            p.style = 'NEAB Heading 2'

            for action in priority_actions[:5]:
                action_text = action if len(action) < 150 else action[:147] + "..."
                p = self.doc.add_paragraph(action_text, style='List Number')
                p.paragraph_format.space_after = Pt(4)

        # Sidbrytning
        self.doc.add_page_break()

    def add_follow_up_plan(self, data):
        """Skapar uppföljningsplan"""
        p = self.doc.add_paragraph('Plan för uppföljning och förbättring')
        p.style = 'NEAB Heading 1'

        plan = data.get('follow_up_plan',
                       'Säkerhetsarbetet följs upp kvartalsvis med rapportering till styrelsen. '
                       'Riskanalys uppdateras årligen eller vid väsentliga förändringar. '
                       'Åtgärder spåras i riskregister med ansvariga och deadlines.')

        p = self.doc.add_paragraph(plan)
        p.style = 'NEAB Body'

    def save(self, filepath):
        """Sparar dokumentet"""
        self.doc.save(filepath)
        return filepath


def create_neab_report(filepath, data):
    """
    Bekvämlighets-funktion för att skapa komplett NEAB Word-rapport

    Args:
        filepath: Sökväg till output DOCX
        data: Dictionary med rapportdata

    Returns:
        Path till genererad DOCX
    """
    template = NEABDocxTemplate()

    # Bygg rapport
    template.add_cover_page(data)
    template.add_executive_summary(data)
    template.add_organization_section(data)
    template.add_regulatory_summary(data)
    template.add_risk_summary(data)
    template.add_mitigation_summary(data)
    template.add_follow_up_plan(data)

    return template.save(filepath)


# Exempel på användning
if __name__ == "__main__":
    example_data = {
        'title': 'NEAB Säkerhetsrapport',
        'subtitle': 'NIS2 och regulatorisk efterlevnad',
        'organization': 'Nordlunda Energi AB',
        'date': '2025-12-15',
        'compliance_rate': 55.6,
        'very_high_risks': 2,
        'high_risks': 5,
        'top_risks': [
            "Otillräcklig segmentering mellan IT och OT",
            "Bristande incidenthanteringsprocess",
            "Saknad formell riskhanteringsprocess",
            "Otillräcklig loggning och övervakning",
            "Bristande backup och återställningsplan"
        ],
        'regulatory_gaps': [
            "NIS2: Styrelseutbildning saknas",
            "GDPR: DPIA för mätdata ej genomförd",
            "Elmarknadslag: Dokumentation av säkerhetsåtgärder ofullständig",
            "NIS2: Incidentrapportering ej implementerad",
            "GDPR: Personuppgiftsbiträdesavtal saknas med vissa leverantörer"
        ],
        'priority_actions': [
            "Implementera Zero Trust-segmentering mellan IT och OT. Införa EDR på alla endpoints.",
            "Komplettera MFA-införande till 100%. Genomföra obligatorisk phishing-utbildning kvartalsvis.",
            "Åtgärda identifierad risk: Bristande efterlevnad av GDPR Art. 35",
            "Införa PAM (Privileged Access Management) för all leverantörsåtkomst med MFA och tidsbegränsning",
            "Implementera API-säkerhetslösning med rate limiting, token rotation och logging"
        ],
        'total_requirements': 9,
        'fulfilled': 2,
        'partial': 6,
        'not_fulfilled': 1,
        'core_business': 'Nordlunda Energi AB är ett regionalt energibolag som driver elnätsverksamhet (DSO) och fjärrvärme. Bolaget betjänar ca 85000 elnätskunder och 12000 fjärrvärmekunder.',
        'critical_systems': 'SCADA/DCS för styrning, AMI för mätdata, OMS för avbrottshantering, samt kund- och faktureringssystem.',
        'security_governance': 'CISO med stab, OT-säkerhetsansvarig, samt SOC via MSSP med 24/7-övervakning.',
        'priority_areas': [
            'Styrelseutbildning (NIS2)',
            'DPIA för mätdata (GDPR)',
            'Formalisering av riskhanteringsprocesser'
        ],
        'risk_distribution': {
            'Mycket hög': 2,
            'Hög': 5,
            'Medel': 8,
            'Låg': 3
        },
        'follow_up_plan': 'Säkerhetsarbetet följs upp kvartalsvis med rapportering till styrelsen. Riskanalys uppdateras årligen eller vid väsentliga förändringar. Åtgärder spåras i riskregister med ansvariga och deadlines.'
    }

    output_path = '/mnt/user-data/outputs/neab_rapport_python_example.docx'
    create_neab_report(output_path, example_data)
    print(f"✓ Rapport skapad: {output_path}")
