"""
Backup Excel-generator för NEAB Security Assessment
Används om RapportGenerator.generera_komplett_paket() kraschar
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
import logging

logger = logging.getLogger(__name__)


class ExcelBackupGenerator:
    """Genererar komplett Excel-rapport från resultat-dict"""

    # Färgschema
    COLORS = {
        'header': 'D32F2F',      # NEAB röd
        'subheader': '1976D2',   # Blå
        'high': 'F44336',        # Röd
        'medium': 'FF9800',      # Orange
        'low': '4CAF50',         # Grön
        'light_gray': 'F5F5F5'
    }

    def __init__(self, output_path):
        self.output_path = Path(output_path)
        self.wb = openpyxl.Workbook()
        # Ta bort standard sheet
        if 'Sheet' in self.wb.sheetnames:
            del self.wb['Sheet']

    def generate(self, resultat):
        """Generera komplett Excel från resultat-dict"""
        try:
            self._create_regulatory_sheet(resultat)
            self._create_risks_sheet(resultat)
            self._create_actions_sheet(resultat)
            self._create_api_sheet(resultat)
            self._create_summary_sheet(resultat)

            self.wb.save(self.output_path)
            logger.info(f"Excel backup skapad: {self.output_path}")
            return self.output_path

        except Exception as e:
            logger.error(f"Excel backup-generering misslyckades: {e}")
            raise

    def _create_api_sheet(self, resultat):
        """API-risker - exakt som original"""
        ws = self.wb.create_sheet("API-risker")

        # Header - exakt som original
        headers = ['Scenario-ID', 'Typ (B2B/B2C)', 'Namn', 'Kryptering',
                   'Autentisering', 'Antal risker', 'Regelverk', 'CIS Controls']

        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color=self.COLORS['subheader'],
                                   end_color=self.COLORS['subheader'], fill_type='solid')
            cell.alignment = Alignment(horizontal='center', wrap_text=True)

        # Data
        api_data = resultat.get('api', {})
        scenarios = api_data.get('scenarios', api_data.get('scenarion', []))

        for row, scenario in enumerate(scenarios, 2):
            if isinstance(scenario, dict):
                ws.cell(row, 1, scenario.get('Scenario-ID', scenario.get('scenario_id', '')))
                typ = scenario.get('Typ', scenario.get('typ', ''))
                ws.cell(row, 2, typ)
                ws.cell(row, 3, scenario.get('Namn', scenario.get('namn', '')))
                ws.cell(row, 4, scenario.get('Kryptering', ''))
                ws.cell(row, 5, scenario.get('Autentisering', ''))

                # Räkna risker
                risker = scenario.get('risker', scenario.get('Risker', []))
                ws.cell(row, 6, len(risker) if isinstance(risker, list) else 0)

                ws.cell(row, 7, scenario.get('Regelverk', ''))
                ws.cell(row, 8, scenario.get('CIS Controls', ''))

        # Justera bredder - exakt som original
        ws.column_dimensions['A'].width = 12   # Scenario-ID
        ws.column_dimensions['B'].width = 25   # Typ
        ws.column_dimensions['C'].width = 40   # Namn
        ws.column_dimensions['D'].width = 35   # Kryptering
        ws.column_dimensions['E'].width = 40   # Autentisering
        ws.column_dimensions['F'].width = 15   # Antal risker
        ws.column_dimensions['G'].width = 30   # Regelverk
        ws.column_dimensions['H'].width = 40   # CIS Controls

    def _create_summary_sheet(self, resultat):
        """Sammanfattning"""
        ws = self.wb.create_sheet("Sammanfattning", 0)

        # Titel
        ws['A1'] = 'NEAB Security Assessment - Sammanfattning'
        ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
        ws['A1'].fill = PatternFill(start_color=self.COLORS['header'],
                                     end_color=self.COLORS['header'],
                                     fill_type='solid')
        ws.merge_cells('A1:D1')

        row = 3

        # Regelefterlevnad
        reg_data = resultat.get('regulatorisk', {})
        krav = reg_data.get('krav', [])
        efterlevnad = reg_data.get('efterlevnad_procent', 0)

        # Om efterlevnad är 0, räkna om från krav
        if efterlevnad == 0 and krav:
            uppfyllt = sum(1 for k in krav if k.get('Status', k.get('status', '')) == 'Uppfyllt')
            efterlevnad = (uppfyllt / len(krav)) * 100

        ws[f'A{row}'] = 'Regelefterlevnad'
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'] = f"{int(round(efterlevnad))}%"  # Heltal istället för decimal

        color = self.COLORS['high'] if efterlevnad < 50 else (
            self.COLORS['medium'] if efterlevnad < 80 else self.COLORS['low'])
        ws[f'B{row}'].fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
        ws[f'B{row}'].font = Font(bold=True, color='FFFFFF')
        row += 1

        # Risker - räkna om från faktiska risker
        risk_data = resultat.get('risk', {})
        risker = risk_data.get('risker', [])

        # Räkna riskfördelning från individuella risker
        actual_distribution = {}
        for risk in risker:
            level = risk.get('Risknivå', 'Okänd')
            actual_distribution[level] = actual_distribution.get(level, 0) + 1

        ws[f'A{row}'] = 'Totalt risker'
        ws[f'B{row}'] = len(risker)
        row += 1

        ws[f'A{row}'] = 'Mycket höga risker'
        ws[f'B{row}'] = actual_distribution.get('Mycket hög', 0)
        ws[f'B{row}'].fill = PatternFill(start_color=self.COLORS['high'],
                                         end_color=self.COLORS['high'], fill_type='solid')
        ws[f'B{row}'].font = Font(bold=True, color='FFFFFF')
        row += 1

        ws[f'A{row}'] = 'Höga risker'
        ws[f'B{row}'] = actual_distribution.get('Hög', 0)
        ws[f'B{row}'].fill = PatternFill(start_color=self.COLORS['medium'],
                                         end_color=self.COLORS['medium'], fill_type='solid')
        ws[f'B{row}'].font = Font(bold=True)
        row += 1

        ws[f'A{row}'] = 'Medelrisker'
        ws[f'B{row}'] = actual_distribution.get('Medel', 0)
        ws[f'B{row}'].fill = PatternFill(start_color='FFEB3B',  # Gul
                                         end_color='FFEB3B', fill_type='solid')
        ws[f'B{row}'].font = Font(bold=True)
        row += 1

        # Åtgärder
        mitigering_data = resultat.get('mitigering', {})
        atgarder = mitigering_data.get('åtgärder', mitigering_data.get('atgarder', []))

        row += 1
        ws[f'A{row}'] = 'Föreslagna åtgärder'
        ws[f'B{row}'] = len(atgarder)

        # Justera kolumnbredd
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 15

    def _create_risks_sheet(self, resultat):
        """Riskmatris - exakt som original"""
        ws = self.wb.create_sheet("Riskmatris")

        # Header - exakt som original
        headers = ['Risk-ID', 'Tillgång / område', 'Hot', 'STRIDE', 'Sårbarhet',
                   'Konsekvens', 'CIA-påverkan', 'Konsekvensnivå', 'Sannolikhet',
                   'Riskvärde', 'Risknivå', 'Befintliga kontroller', 'Regelverk', 'Kommentar']

        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color=self.COLORS['subheader'],
                                   end_color=self.COLORS['subheader'], fill_type='solid')
            cell.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)

        # Data
        risk_data = resultat.get('risk', {})
        risker = risk_data.get('risker', [])

        for row, risk in enumerate(risker, 2):
            # Extrahera alla fält med fallback
            ws.cell(row, 1, risk.get('Risk-ID', ''))
            ws.cell(row, 2, risk.get('Tillgång / område', risk.get('Tillgång', '')))
            ws.cell(row, 3, risk.get('Hot', ''))
            ws.cell(row, 4, risk.get('STRIDE-kategori', risk.get('STRIDE', '')))
            ws.cell(row, 5, risk.get('Sårbarhet', ''))
            ws.cell(row, 6, risk.get('Konsekvens', ''))
            ws.cell(row, 7, risk.get('CIA-påverkan', ''))
            ws.cell(row, 8, risk.get('Konsekvensnivå', risk.get('Konsekvens_niveau', '')))
            ws.cell(row, 9, risk.get('Sannolikhet', ''))
            ws.cell(row, 10, risk.get('Riskvärde', ''))

            # Risknivå med färgkodning
            riskniva = risk.get('Risknivå', '')
            cell = ws.cell(row, 11, riskniva)
            if riskniva == 'Mycket hög':
                cell.fill = PatternFill(start_color=self.COLORS['high'],
                                       end_color=self.COLORS['high'], fill_type='solid')
                cell.font = Font(color='FFFFFF', bold=True)
            elif riskniva == 'Hög':
                cell.fill = PatternFill(start_color=self.COLORS['medium'],
                                       end_color=self.COLORS['medium'], fill_type='solid')
                cell.font = Font(bold=True)
            elif riskniva == 'Medel':
                cell.fill = PatternFill(start_color='FFEB3B',  # Gul
                                       end_color='FFEB3B', fill_type='solid')
                cell.font = Font(bold=True)
            elif riskniva == 'Låg':
                cell.fill = PatternFill(start_color=self.COLORS['low'],
                                       end_color=self.COLORS['low'], fill_type='solid')
                cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

            ws.cell(row, 12, risk.get('Befintliga kontroller', ''))
            ws.cell(row, 13, risk.get('Regelverk', ''))
            ws.cell(row, 14, risk.get('Kommentar', ''))

            # Alignment för alla celler i raden
            for col in range(1, 15):
                cell = ws.cell(row, col)
                if col in [1, 7, 8, 9, 10, 11]:  # Numeriska/korta kolumner
                    cell.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
                else:  # Textkolumner
                    cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

        # Sätt radhöjd för header
        ws.row_dimensions[1].height = 30

        # Justera bredder - exakt som original
        ws.column_dimensions['A'].width = 10   # Risk-ID
        ws.column_dimensions['B'].width = 35   # Tillgång
        ws.column_dimensions['C'].width = 30   # Hot
        ws.column_dimensions['D'].width = 25   # STRIDE
        ws.column_dimensions['E'].width = 40   # Sårbarhet
        ws.column_dimensions['F'].width = 35   # Konsekvens
        ws.column_dimensions['G'].width = 12   # CIA
        ws.column_dimensions['H'].width = 15   # Konsekvensnivå
        ws.column_dimensions['I'].width = 12   # Sannolikhet
        ws.column_dimensions['J'].width = 12   # Riskvärde
        ws.column_dimensions['K'].width = 15   # Risknivå
        ws.column_dimensions['L'].width = 50   # Befintliga kontroller
        ws.column_dimensions['M'].width = 25   # Regelverk
        ws.column_dimensions['N'].width = 40   # Kommentar

    def _create_regulatory_sheet(self, resultat):
        """Regulatoriska krav - exakt som original"""
        ws = self.wb.create_sheet("Regulatorisk kartläggning")

        # Header - exakt som original
        headers = ['Regelverk', 'Artikel', 'Krav', 'Tillämpning', 'Status',
                   'Gap-analys', 'Föreslagen åtgärd', 'CIS Controls']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color=self.COLORS['subheader'],
                                   end_color=self.COLORS['subheader'], fill_type='solid')

        # Data
        reg_data = resultat.get('regulatorisk', {})
        krav = reg_data.get('krav', [])

        for row, krav_item in enumerate(krav, 2):
            # Extrahera alla fält med fallback till olika namnkonventioner
            regelverk = krav_item.get('Regelverk', krav_item.get('regelverk', ''))
            artikel = krav_item.get('Artikel', krav_item.get('artikel', ''))
            krav_text = krav_item.get('Krav', krav_item.get('krav_beskrivning', ''))
            tillampning = krav_item.get('Tillämpning', krav_item.get('tillampning', ''))
            status = krav_item.get('Status', krav_item.get('status', ''))
            gap = krav_item.get('Gap-analys', krav_item.get('gap_analys', krav_item.get('Kommentar', '')))
            atgard = krav_item.get('Föreslagen åtgärd', krav_item.get('foreslagen_atgard', ''))
            cis = krav_item.get('CIS Controls', krav_item.get('cis_controls', ''))

            ws.cell(row, 1, regelverk)
            ws.cell(row, 2, artikel)
            ws.cell(row, 3, krav_text)
            ws.cell(row, 4, tillampning)

            # Status med färgkodning
            cell = ws.cell(row, 5, status)
            if status == 'Uppfyllt':
                cell.fill = PatternFill(start_color=self.COLORS['low'],
                                       end_color=self.COLORS['low'], fill_type='solid')
            elif status == 'Delvis uppfyllt':
                cell.fill = PatternFill(start_color=self.COLORS['medium'],
                                       end_color=self.COLORS['medium'], fill_type='solid')
            elif status == 'Ej uppfyllt':
                cell.fill = PatternFill(start_color=self.COLORS['high'],
                                       end_color=self.COLORS['high'], fill_type='solid')
                cell.font = Font(color='FFFFFF')

            ws.cell(row, 6, gap)
            ws.cell(row, 7, atgard)
            ws.cell(row, 8, cis)

            # Text wrapping för alla celler
            for col in range(1, 9):
                cell = ws.cell(row, col)
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)

        # Justera bredder - bredare än original för bättre läsbarhet
        ws.column_dimensions['A'].width = 15   # Regelverk (var 12)
        ws.column_dimensions['B'].width = 12   # Artikel (var 10)
        ws.column_dimensions['C'].width = 55   # Krav (var 45)
        ws.column_dimensions['D'].width = 35   # Tillämpning (var 30)
        ws.column_dimensions['E'].width = 20   # Status (var 18)
        ws.column_dimensions['F'].width = 60   # Gap-analys (var 50)
        ws.column_dimensions['G'].width = 60   # Föreslagen åtgärd (var 50)
        ws.column_dimensions['H'].width = 30   # CIS Controls (var 25)

    def _create_actions_sheet(self, resultat):
        """Riskbehandlingsplan - exakt som original"""
        ws = self.wb.create_sheet("Riskbehandlingsplan")

        # Header - exakt som original
        headers = ['Åtgärds-ID', 'Risk-ID', 'Beskrivning', 'Strategi', 'CIS Controls',
                   'Ansvarig roll', 'Planerat datum', 'Status', 'Kostnad', 'Effekt',
                   'Kvarvarande risk', 'Implementation']

        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(start_color=self.COLORS['subheader'],
                                   end_color=self.COLORS['subheader'], fill_type='solid')
            cell.alignment = Alignment(horizontal='center', wrap_text=True)

        # Data
        mitigering_data = resultat.get('mitigering', {})
        atgarder = mitigering_data.get('åtgärder', mitigering_data.get('atgarder', []))

        for row, atgard in enumerate(atgarder, 2):
            if isinstance(atgard, dict):
                ws.cell(row, 1, atgard.get('Åtgärds-ID', atgard.get('atgard_id', '')))
                ws.cell(row, 2, atgard.get('Risk-ID', ''))
                ws.cell(row, 3, atgard.get('Beskrivning', atgard.get('beskrivning', '')))
                ws.cell(row, 4, atgard.get('Strategi', 'Reducera'))
                ws.cell(row, 5, atgard.get('CIS Controls', atgard.get('cis_controls', '')))
                ws.cell(row, 6, atgard.get('Ansvarig roll', atgard.get('ansvarig', '')))
                ws.cell(row, 7, atgard.get('Planerat datum', ''))
                ws.cell(row, 8, atgard.get('Status', 'Planerad'))
                ws.cell(row, 9, atgard.get('Kostnad', 'M'))
                ws.cell(row, 10, atgard.get('Effekt', 'M'))

                # Kvarvarande risk med färgkodning
                kvarvarande = atgard.get('Kvarvarande risk', 'Medel')
                cell = ws.cell(row, 11, kvarvarande)
                if kvarvarande == 'Mycket hög':
                    cell.fill = PatternFill(start_color=self.COLORS['high'],
                                           end_color=self.COLORS['high'], fill_type='solid')
                    cell.font = Font(color='FFFFFF', bold=True)
                elif kvarvarande == 'Hög':
                    cell.fill = PatternFill(start_color=self.COLORS['medium'],
                                           end_color=self.COLORS['medium'], fill_type='solid')
                    cell.font = Font(bold=True)
                elif kvarvarande == 'Medel':
                    cell.fill = PatternFill(start_color='FFEB3B',  # Gul
                                           end_color='FFEB3B', fill_type='solid')
                    cell.font = Font(bold=True)
                elif kvarvarande == 'Låg':
                    cell.fill = PatternFill(start_color=self.COLORS['low'],
                                           end_color=self.COLORS['low'], fill_type='solid')
                    cell.font = Font(bold=True)

                ws.cell(row, 12, atgard.get('Implementation', ''))

                # Text wrapping för textkolumner
                for col in [3, 5, 6, 12]:  # Beskrivning, CIS, Ansvarig, Implementation
                    ws.cell(row, col).alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
            else:
                # Fallback om åtgärd bara är en sträng
                ws.cell(row, 1, f"Å{row-1:03d}")
                ws.cell(row, 3, str(atgard))
                ws.cell(row, 4, 'Reducera')
                ws.cell(row, 8, 'Planerad')

        # Justera bredder - bredare för bättre läsbarhet
        ws.column_dimensions['A'].width = 12   # Åtgärds-ID
        ws.column_dimensions['B'].width = 10   # Risk-ID
        ws.column_dimensions['C'].width = 70   # Beskrivning (var 60)
        ws.column_dimensions['D'].width = 12   # Strategi
        ws.column_dimensions['E'].width = 30   # CIS Controls (var 25)
        ws.column_dimensions['F'].width = 35   # Ansvarig roll (var 30)
        ws.column_dimensions['G'].width = 15   # Planerat datum
        ws.column_dimensions['H'].width = 12   # Status
        ws.column_dimensions['I'].width = 10   # Kostnad
        ws.column_dimensions['J'].width = 10   # Effekt
        ws.column_dimensions['K'].width = 20   # Kvarvarande risk (var 18)
        ws.column_dimensions['L'].width = 70   # Implementation (var 60)


def generate_excel_backup(resultat, output_path):
    """
    Convenience function för att generera backup Excel

    Args:
        resultat: Resultat-dict från analysmodulerna
        output_path: Var Excel-filen ska sparas

    Returns:
        Path till genererad Excel-fil
    """
    generator = ExcelBackupGenerator(output_path)
    return generator.generate(resultat)