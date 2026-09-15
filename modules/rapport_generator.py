"""
Rapport Generator
=================
Genererar Excel-arbetsböcker och PDF-rapporter enligt gruppexaminationens format

Del av: NEAB Security Assessment Suite
"""

import logging
from pathlib import Path
from typing import Dict, Any, List
from datetime import datetime
import json

# Excel-bibliotek
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logging.warning("openpyxl inte installerat - Excel-funktionalitet begränsad")

# PDF-bibliotek
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER

    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False
    logging.warning("reportlab inte installerat - PDF-funktionalitet begränsad")

logger = logging.getLogger(__name__)


class RapportGenerator:
    """
    Genererar rapporter och Excel-dokument
    """

    def __init__(self, config: Dict[str, Any], output_dir: Path):
        self.config = config
        self.output_dir = output_dir
        self.organisation = config.get("organisation", {})

        logger.info(f"RapportGenerator initierad - Output: {output_dir}")

    def generera_komplett_paket(self, resultat: Dict[str, Any]) -> Dict[str, Path]:
        """
        Genererar alla outputs enligt gruppexaminationens krav

        Returns:
            Dictionary med sökvägar till genererade filer
        """
        logger.info("Genererar komplett rapportpaket...")

        filer = {}

        # 1. Excel-arbetsbok med alla flikar
        if OPENPYXL_AVAILABLE:
            excel_path = self.skapa_excel_workbook(resultat)
            filer["excel"] = excel_path
            logger.info(f"✓ Excel skapad: {excel_path.name}")
        else:
            logger.warning("Excel kan ej genereras - openpyxl saknas")

        # 2. PDF-ledningsrapport
        if REPORTLAB_AVAILABLE:
            pdf_path = self.skapa_ledningsrapport(resultat)
            filer["pdf"] = pdf_path
            logger.info(f"✓ PDF skapad: {pdf_path.name}")
        else:
            logger.warning("PDF kan ej genereras - reportlab saknas")

        # 3. JSON-export (alltid tillgänglig)
        json_path = self.exportera_json(resultat)
        filer["json"] = json_path
        logger.info(f"✓ JSON skapad: {json_path.name}")

        # 4. Workshop-protokoll (simulerad för nu)
        protokoll_path = self.skapa_workshop_protokoll(resultat)
        filer["protokoll"] = protokoll_path
        logger.info(f"✓ Protokoll skapat: {protokoll_path.name}")

        logger.info("Komplett rapportpaket genererat!")
        return filer

    def skapa_excel_workbook(self, resultat: Dict[str, Any]) -> Path:
        """
        Skapar Excel-arbetsbok med alla flikar enligt kraven:
        - Regulatorisk kartläggning
        - Riskmatris
        - Riskbehandlingsplan
        - API-risker
        """
        if not OPENPYXL_AVAILABLE:
            logger.error("Kan inte skapa Excel - openpyxl inte tillgängligt")
            return None

        wb = Workbook()
        wb.remove(wb.active)  # Ta bort default sheet

        # Flik 1: Regulatorisk kartläggning
        self._skapa_regulatorisk_flik(wb, resultat.get("regulatorisk_kartläggning", {}))

        # Flik 2: Riskmatris
        self._skapa_riskmatris_flik(wb, resultat.get("riskanalys", {}))

        # Flik 3: Riskbehandlingsplan
        self._skapa_riskbehandling_flik(wb, resultat.get("mitigeringar", {}))

        # Flik 4: API-risker
        self._skapa_api_flik(wb, resultat.get("api_analys", {}))

        # Spara
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_path = self.output_dir / f"NEAB_Sakerhetsanalys_{timestamp}.xlsx"
        wb.save(excel_path)

        return excel_path

    def _skapa_regulatorisk_flik(self, wb: Workbook, data: Dict):
        """Skapar flik: Regulatorisk kartläggning"""
        ws = wb.create_sheet("Regulatorisk kartläggning")

        # Headers
        headers = ["Regelverk", "Artikel", "Krav", "Tillämpning", "Status",
                   "Gap-analys", "Föreslagen åtgärd", "CIS Controls"]

        # Stil för headers
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data
        krav_lista = data.get("krav", [])
        for row_num, krav in enumerate(krav_lista, 2):
            ws.cell(row=row_num, column=1).value = krav.get("Regelverk")
            ws.cell(row=row_num, column=2).value = krav.get("Artikel")
            ws.cell(row=row_num, column=3).value = krav.get("Krav")
            ws.cell(row=row_num, column=4).value = krav.get("Tillämpning")

            # Status med färgkodning
            status_cell = ws.cell(row=row_num, column=5)
            status_cell.value = krav.get("Status")
            if krav.get("Status") == "Uppfyllt":
                status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            elif krav.get("Status") == "Delvis uppfyllt":
                status_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            else:
                status_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

            ws.cell(row=row_num, column=6).value = krav.get("Gap-analys")
            ws.cell(row=row_num, column=7).value = krav.get("Föreslagen åtgärd")
            ws.cell(row=row_num, column=8).value = krav.get("CIS Controls")

        # Autowidth
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Text wrap för långa celler
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    def _skapa_riskmatris_flik(self, wb: Workbook, data: Dict):
        """Skapar flik: Riskmatris"""
        ws = wb.create_sheet("Riskmatris")

        # Headers
        headers = ["Risk-ID", "Tillgång / område", "Hot", "STRIDE", "Sårbarhet",
                   "Konsekvens", "CIA-påverkan", "Konsekvensnivå", "Sannolikhet",
                   "Riskvärde", "Risknivå", "Befintliga kontroller", "Regelverk", "Kommentar"]

        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data
        risker = data.get("risker", [])
        for row_num, risk in enumerate(risker, 2):
            ws.cell(row=row_num, column=1).value = risk.get("Risk-ID")
            ws.cell(row=row_num, column=2).value = risk.get("Tillgång / område")
            ws.cell(row=row_num, column=3).value = risk.get("Hot")
            ws.cell(row=row_num, column=4).value = risk.get("STRIDE")
            ws.cell(row=row_num, column=5).value = risk.get("Sårbarhet")
            ws.cell(row=row_num, column=6).value = risk.get("Konsekvens")
            ws.cell(row=row_num, column=7).value = risk.get("CIA-påverkan")
            ws.cell(row=row_num, column=8).value = risk.get("Konsekvensnivå")
            ws.cell(row=row_num, column=9).value = risk.get("Sannolikhet")
            ws.cell(row=row_num, column=10).value = risk.get("Riskvärde")

            # Risknivå med färgkodning
            riskniva_cell = ws.cell(row=row_num, column=11)
            riskniva_cell.value = risk.get("Risknivå")
            if risk.get("Risknivå") == "Mycket hög":
                riskniva_cell.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
                riskniva_cell.font = Font(color="FFFFFF", bold=True)
            elif risk.get("Risknivå") == "Hög":
                riskniva_cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            elif risk.get("Risknivå") == "Medel":
                riskniva_cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            else:
                riskniva_cell.fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")

            ws.cell(row=row_num, column=12).value = risk.get("Befintliga kontroller")
            ws.cell(row=row_num, column=13).value = risk.get("Regelverk")
            ws.cell(row=row_num, column=14).value = risk.get("Kommentar")

        # Autowidth och wrap
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    def _skapa_riskbehandling_flik(self, wb: Workbook, data: Dict):
        """Skapar flik: Riskbehandlingsplan"""
        ws = wb.create_sheet("Riskbehandlingsplan")

        headers = ["Åtgärds-ID", "Risk-ID", "Beskrivning", "Strategi", "CIS Controls",
                   "Ansvarig roll", "Planerat datum", "Status", "Kostnad", "Effekt",
                   "Kvarvarande risk", "Implementation"]

        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        åtgärder = data.get("åtgärder", [])
        for row_num, åtgärd in enumerate(åtgärder, 2):
            ws.cell(row=row_num, column=1).value = åtgärd.get("Åtgärds-ID")
            ws.cell(row=row_num, column=2).value = åtgärd.get("Risk-ID")
            ws.cell(row=row_num, column=3).value = åtgärd.get("Beskrivning")
            ws.cell(row=row_num, column=4).value = åtgärd.get("Strategi")
            ws.cell(row=row_num, column=5).value = åtgärd.get("CIS Controls")
            ws.cell(row=row_num, column=6).value = åtgärd.get("Ansvarig roll")
            ws.cell(row=row_num, column=7).value = åtgärd.get("Planerat datum")
            ws.cell(row=row_num, column=8).value = åtgärd.get("Status")
            ws.cell(row=row_num, column=9).value = åtgärd.get("Kostnad")
            ws.cell(row=row_num, column=10).value = åtgärd.get("Effekt")
            ws.cell(row=row_num, column=11).value = åtgärd.get("Kvarvarande risk")
            ws.cell(row=row_num, column=12).value = åtgärd.get("Implementation")

        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    def _skapa_api_flik(self, wb: Workbook, data: Dict):
        """Skapar flik: API-risker"""
        ws = wb.create_sheet("API-risker")

        headers = ["Scenario-ID", "Typ (B2B/B2C)", "Namn", "Kryptering",
                   "Autentisering", "Antal risker", "Regelverk", "CIS Controls"]

        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        scenarios = data.get("scenarios", [])
        for row_num, scenario in enumerate(scenarios, 2):
            ws.cell(row=row_num, column=1).value = scenario.get("Scenario-ID")
            ws.cell(row=row_num, column=2).value = scenario.get("Typ")
            ws.cell(row=row_num, column=3).value = scenario.get("Namn")
            ws.cell(row=row_num, column=4).value = scenario.get("Kryptering")
            ws.cell(row=row_num, column=5).value = scenario.get("Autentisering")
            ws.cell(row=row_num, column=6).value = scenario.get("Antal_risker")
            ws.cell(row=row_num, column=7).value = scenario.get("Regelverk")
            ws.cell(row=row_num, column=8).value = scenario.get("CIS_Controls")

        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            adjusted_width = min(max_length + 2, 40)
            ws.column_dimensions[column_letter].width = adjusted_width

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    def skapa_ledningsrapport(self, resultat: Dict[str, Any]) -> Path:
        """
        Skapar PDF-ledningsrapport enligt struktur i gruppexamination
        """
        if not REPORTLAB_AVAILABLE:
            logger.error("Kan inte skapa PDF - reportlab inte tillgängligt")
            # Skapa enkel textfil istället
            return self._skapa_text_rapport(resultat)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        pdf_path = self.output_dir / f"NEAB_Ledningsrapport_{timestamp}.pdf"

        doc = SimpleDocTemplate(str(pdf_path), pagesize=A4,
                                topMargin=2 * cm, bottomMargin=2 * cm)

        story = []
        styles = getSampleStyleSheet()

        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#366092'),
            spaceAfter=30,
            alignment=TA_CENTER
        )

        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor('#366092'),
            spaceAfter=12,
            spaceBefore=12
        )

        # Titel
        story.append(Paragraph("Riskworkshop med regulatorisk kartläggning", title_style))
        story.append(Paragraph(f"{self.organisation.get('namn', 'Nordlunda Energi AB')}", styles['Heading2']))
        story.append(Spacer(1, 1 * cm))

        # Sammanfattning för ledningen
        story.append(Paragraph("Sammanfattning för ledningsnivå", heading_style))
        sammanfattning_text = self._generera_ledningssammanfattning(resultat)
        story.append(Paragraph(sammanfattning_text, styles['BodyText']))
        story.append(Spacer(1, 0.5 * cm))

        # Organisationsbeskrivning
        story.append(Paragraph("Organisationen och kritiska processer", heading_style))
        org_text = self._generera_org_beskrivning()
        story.append(Paragraph(org_text, styles['BodyText']))
        story.append(Spacer(1, 0.5 * cm))

        # Regulatorisk sammanfattning
        story.append(Paragraph("Regulatorisk kartläggning", heading_style))
        reg_data = resultat.get("regulatorisk_kartläggning", {})
        reg_text = reg_data.get("sammanfattning", "Ingen data tillgänglig")
        story.append(Paragraph(reg_text, styles['BodyText']))
        story.append(Spacer(1, 0.5 * cm))

        # Riskanalyssammanfattning
        story.append(Paragraph("Sammanfattning av riskanalys", heading_style))
        risk_data = resultat.get("riskanalys", {})
        risk_text = risk_data.get("sammanfattning", "Ingen data tillgänglig")
        story.append(Paragraph(risk_text, styles['BodyText']))
        story.append(Spacer(1, 0.5 * cm))

        # Mitigeringsstrategier
        story.append(Paragraph("Mitigeringsstrategier och rekommenderade åtgärder", heading_style))
        mit_data = resultat.get("mitigeringar", {})
        mit_text = mit_data.get("sammanfattning", "Ingen data tillgänglig")
        story.append(Paragraph(mit_text, styles['BodyText']))
        story.append(Spacer(1, 0.5 * cm))

        # Uppföljningsplan
        story.append(Paragraph("Plan för uppföljning och förbättring", heading_style))
        uppfoljning_text = """
Säkerhetsarbetet följs upp kvartalsvis med rapportering till styrelsen. 
Riskanalys uppdateras årligen eller vid väsentliga förändringar. 
Åtgärder spåras i riskregister med ansvariga och deadlines.
        """
        story.append(Paragraph(uppfoljning_text, styles['BodyText']))

        # Bygg PDF
        doc.build(story)

        return pdf_path

    def _skapa_text_rapport(self, resultat: Dict[str, Any]) -> Path:
        """Skapar textbaserad rapport om PDF ej är tillgängligt"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        txt_path = self.output_dir / f"NEAB_Ledningsrapport_{timestamp}.txt"

        with open(txt_path, 'w', encoding='utf-8') as f:
            f.write("=" * 80 + "\n")
            f.write("RISKWORKSHOP MED REGULATORISK KARTLÄGGNING\n")
            f.write(f"{self.organisation.get('namn', 'Nordlunda Energi AB')}\n")
            f.write("=" * 80 + "\n\n")

            f.write("SAMMANFATTNING FÖR LEDNINGEN\n")
            f.write("-" * 80 + "\n")
            f.write(self._generera_ledningssammanfattning(resultat))
            f.write("\n\n")

            # Lägg till övriga sektioner...

        return txt_path

    def _generera_ledningssammanfattning(self, resultat: Dict[str, Any]) -> str:
        """Genererar koncis sammanfattning för ledningen"""
        reg_data = resultat.get("regulatorisk_kartläggning", {})
        risk_data = resultat.get("riskanalys", {})

        reg_stats = reg_data.get("statistik", {})

        risker = risk_data.get("risker", [])
        mycket_hog = sum(1 for r in risker if r.get("Risknivå") == "Mycket hög")
        hog = sum(1 for r in risker if r.get("Risknivå") == "Hög")

        return f"""
<b>Övergripande bedömning:</b> NEAB har en {reg_stats.get('uppfyllnadsgrad_procent', 0)}% 
regelefterlevnad med {mycket_hog} mycket höga och {hog} höga risker identifierade.

<b>Viktigaste riskområden:</b>
• Ransomware med risk för sidoförflyttning IT→OT
• Leverantörskedjeangrepp via fjärrsupport
• API-säkerhet (särskilt kundportaler)

<b>Regulatoriska brister:</b>
• Styrelseutbildning i cybersäkerhet (NIS2)
• DPIA för AMI-system (GDPR)
• Dokumentation av fysisk säkerhet (CER)

<b>Föreslagna prioriterade åtgärder:</b>
1. Implementera Zero Trust-segmentering IT/OT
2. Inför PAM för leverantörsåtkomst
3. Genomför styrelseutbildning och DPIA
        """.strip()

    def _generera_org_beskrivning(self) -> str:
        """Genererar organisationsbeskrivning"""
        return f"""
<b>Kärnverksamhet:</b> {self.organisation.get('namn', 'Nordlunda Energi AB')} är ett regionalt 
energibolag som driver elnätsverksamhet (DSO) och fjärrvärme. Bolaget betjänar ca 
{self.organisation.get('antal_kunder_el', 85000)} elnätskunder och 
{self.organisation.get('antal_kunder_fjarr', 12000)} fjärrvärmekunder.

<b>Kritiska system:</b> SCADA/DCS för styrning, AMI för mätdata, OMS för avbrottshantering, 
samt kund- och faktureringssystem.

<b>Säkerhetsstyrning:</b> CISO med stab, OT-säkerhetsansvarig, samt SOC via MSSP med 
24/7-övervakning.
        """.strip()

    def skapa_workshop_protokoll(self, resultat: Dict[str, Any]) -> Path:
        """Skapar workshopprotokoll (simulerat för nu)"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        protokoll_path = self.output_dir / f"Workshop_Protokoll_{timestamp}.txt"

        with open(protokoll_path, 'w', encoding='utf-8') as f:
            f.write("RISKWORKSHOP - PROTOKOLL\n")
            f.write("=" * 80 + "\n")
            f.write(f"Datum: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n")
            f.write(f"Organisation: {self.organisation.get('namn')}\n")
            f.write(f"Deltagare: CISO, IT-chef, Driftchef, OT-säkerhetsansvarig\n")
            f.write("\n")
            f.write("AGENDA:\n")
            f.write("1. Genomgång av regulatoriska krav\n")
            f.write("2. Identifiering av tillgångar och beroenden\n")
            f.write("3. Hotmodellering med STRIDE\n")
            f.write("4. Konsekvens- och sannolikhetsbedömning\n")
            f.write("5. Rangordning av risker\n")
            f.write("6. Diskussion av mitigeringsstrategier\n")
            f.write("\n")
            f.write("RESULTAT:\n")
            f.write(f"- Identifierade {len(resultat.get('riskanalys', {}).get('risker', []))} risker\n")
            f.write(
                f"- Kartlagda {len(resultat.get('regulatorisk_kartläggning', {}).get('krav', []))} regulatoriska krav\n")
            f.write(f"- Föreslagna {len(resultat.get('mitigeringar', {}).get('åtgärder', []))} åtgärder\n")
            f.write("\n")
            f.write("Workshopen genomfördes enligt MSB:s metodstöd med STRIDE för hotmodellering.\n")

        return protokoll_path

    def exportera_json(self, resultat: Dict[str, Any]) -> Path:
        """Exporterar resultat som JSON"""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        json_path = self.output_dir / f"NEAB_Resultat_{timestamp}.json"

        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(resultat, f, ensure_ascii=False, indent=2)

        return json_path